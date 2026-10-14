"""Render the benchmark case from bible.yaml + source/ into docs/ and expected/.

    uv run python cases/harbor-dispute/generate.py [--force]

docs/ is committed; regeneration is deliberate (bump case_version in bible.yaml).
"""

from __future__ import annotations

import argparse
import hashlib
import io
import random
import re
import shutil
import sys
from datetime import date
from email.message import EmailMessage
from email.utils import format_datetime, parsedate_to_datetime
from pathlib import Path

import yaml

from dossier_bench.render import (  # noqa: E402  (shared renderers; run from the repo root with `uv run python`)
    blocks_to_text,
    canonicalize,
    expand_long_document,
    logo_png,
    parse_md,
    pdf_page_count,
    render_csv,
    render_docx,
    render_eml,
    render_pdf,
    render_png_scan,
    render_scanned_pdf,
    render_sign_photo,
    sha256,
)

HERE = Path(__file__).resolve().parent
SRC = HERE / "source"
DOCS = HERE / "docs"
EXPECTED = HERE / "expected"
ATTACH_ONLY = DOCS / "attachments-only"   # email attachments; never uploaded directly

# ----------------------------------------------------------------------------
# markdown subset → blocks
# ----------------------------------------------------------------------------


# ----------------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------------







# ----------------------------------------------------------------------------
# DOCX / TXT / XLSX / EML / JPG
# ----------------------------------------------------------------------------




LONG_TITLES = ["Electrical installation", "Control system software", "Mechanical installation", "Safety and guarding", "Testing and commissioning", "Documentation and training", "Spare parts and maintenance", "Interfaces to the warehouse management system", "Quality assurance", "Environmental and noise", "Marking and identification", "Packaging, transport and storage", "Site works and access", "Acceptance criteria", "Warranty and defects"]
LONG_SUBJECTS = ["sorter carriage", "induction conveyor", "main control panel", "field network", "photo-eye array", "drive inverter", "belt conveyor", "outfeed chute", "safety fence", "emergency stop circuit", "HMI workstation", "barcode scanner tunnel"]
LONG_REQUIREMENTS = ["be rated for continuous duty at the design throughput", "be supplied with as-built drawings before site acceptance", "be tested over 72 hours of continuous operation", "be protected to IP54 or better", "be labelled with the tag number shown on the schematic", "be interlocked with the adjacent zone", "achieve a mean time between failures of not less than 4,000 hours", "be accessible for maintenance without the removal of guarding"]
LONG_STANDARDS = ["EN ISO 13857", "EN 60204-1", "EN ISO 13849-1", "IEC 61131-3", "EN 619", "DIN 18202", "EN 60529"]


def render_xlsx(bible: dict, out: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Change Orders"
    ws.append(["Project", "Valencia Distribution Centre — Automated Conveyor and Sortation System (NWL-VAL-2024-07)"])
    ws.append(["Register maintained by", "Tidewell Fabrication Ltd (clause 5.5)"])
    ws.append([])
    hdr = ["CO ref", "Date", "Title", "Value (EUR)", "Time (days)", "Status", "Countersigned by Buyer", "Notes"]
    ws.append(hdr)
    for c in ws[4]:
        c.font = Font(bold=True)
    notes = {
        "CO-1": "Second sorter outfeed on recirculation line. Reflected in Amendment No. 1.",
        "CO-2": "Redundant S7-1500 PLC pair. Reflected in Amendment No. 1.",
        "CO-3": "Re-level Bay 3 slab C4–C7 (18 mm over 3 m vs 5 mm tolerance, Levantina LT-2025-0218). Buyer replied 24 Feb 2025: proceed at risk. Not countersigned.",
        "CO-4": "Additional fencing at outfeed lanes 20–26 at Buyer request.",
    }
    signed = {"CO-1": "Yes — L. Ortega 19/12/2024", "CO-2": "Yes — L. Ortega 10/01/2025", "CO-3": "No", "CO-4": "Yes — L. Ortega 28/03/2025"}
    for co in bible["change_orders"]:
        ws.append([co["id"], co["date"], co["title"], co["value_eur"], co["days"], co["status"], signed[co["id"]], notes[co["id"]]])
    ws.append([])
    approved = sum(c["value_eur"] for c in bible["change_orders"] if c["status"] == "approved")
    disputed = sum(c["value_eur"] for c in bible["change_orders"] if c["status"] == "disputed")
    ws.append(["", "", "Total approved", approved, sum(c["days"] for c in bible["change_orders"] if c["status"] == "approved")])
    ws.append(["", "", "Total disputed", disputed, sum(c["days"] for c in bible["change_orders"] if c["status"] == "disputed")])
    ws.append(["", "", "Original contract price", bible["contract"]["price_v1_eur"]])
    ws.append(["", "", "Price after Amendment No. 1", bible["contract"]["price_v2_eur"]])
    for col, w in zip("ABCDEFGH", [10, 12, 44, 14, 12, 12, 26, 70]):
        ws.column_dimensions[col].width = w

    ms = wb.create_sheet("Milestones")
    ms.append(["Milestone", "Baseline v1 (4 Nov 2024)", "Re-baseline v2 (20 Jan 2025)", "Actual / forecast", "Variance vs v2 (days)", "Comment"])
    for c in ms[1]:
        c.font = Font(bold=True)
    v1, v2 = bible["contract"]["milestones_v1"], bible["contract"]["milestones_v2"]
    actual = {
        "design_approval": (date(2024, 12, 13), "Achieved"),
        "delivery": (date(2025, 2, 20), "Last consignment 20 Feb 2025 (sub-supplier)"),
        "mechanical_completion": (date(2025, 3, 21), "Notified 21 Mar 2025 excluding Bay 3 zone C4–C7"),
        "practical_completion": (date(2025, 5, 29), "Forecast per progress report 11; agreement terminated 2 May 2025"),
    }
    for k, label in [("design_approval", "Design Approval"), ("delivery", "Delivery"), ("mechanical_completion", "Mechanical Completion"), ("practical_completion", "Practical Completion")]:
        a, comment = actual[k]
        ms.append([label, v1[k], v2[k], a, (a - v2[k]).days, comment])
    for col, w in zip("ABCDEF", [24, 24, 26, 18, 20, 60]):
        ms.column_dimensions[col].width = w
    from datetime import datetime

    wb.properties.created = wb.properties.modified = datetime(2025, 3, 28, 9, 0, 0)
    wb.save(str(out))


def person(bible: dict, ref: str) -> dict:
    return next(p for p in bible["people"] if p["key"] == ref)


def addr(bible: dict, ref: str) -> str:
    p = person(bible, ref)
    return f'"{p["name"]}" <{p["email"]}>'


def render_eml_thread(bible: dict, emails: dict, attachments: dict[str, Path]) -> list[tuple[str, Path, str]]:
    """Returns [(message key, path, plain body as a human would read it)]."""
    out = []
    by_key = {m["key"]: m for m in emails["messages"]}
    rendered_bodies: dict[str, str] = {}
    for i, m in enumerate(emails["messages"], start=1):
        msg = EmailMessage()
        subject = emails["thread_subject"] if not m.get("in_reply_to") else "RE: " + emails["thread_subject"]
        msg["Subject"] = subject
        msg["From"] = addr(bible, m["from"]["ref"])
        msg["To"] = ", ".join(addr(bible, t["ref"]) for t in m["to"])
        if m.get("cc"):
            msg["Cc"] = ", ".join(addr(bible, t["ref"]) for t in m["cc"])
        msg["Date"] = m["date"]
        msg["Message-ID"] = m["message_id"]
        body = m["body"].rstrip() + "\n"
        if m.get("in_reply_to"):
            parent = by_key[m["in_reply_to"]]
            msg["In-Reply-To"] = parent["message_id"]
            refs = []
            p = parent
            while p is not None:
                refs.insert(0, p["message_id"])
                p = by_key.get(p.get("in_reply_to")) if p.get("in_reply_to") else None
            msg["References"] = " ".join(refs)
            pdt = parsedate_to_datetime(parent["date"])
            quoted = "\n".join("> " + ln for ln in rendered_bodies[parent["key"]].splitlines())
            body = body + f"\n\nOn {pdt.strftime('%a, %d %b %Y at %H:%M')}, {person(bible, parent['from']['ref'])['name']} wrote:\n\n{quoted}\n"
        rendered_bodies[m["key"]] = body
        msg.set_content(body)
        for att_key in m.get("attachments", []):
            path = attachments[att_key]
            msg.add_attachment(path.read_bytes(), maintype="application", subtype="pdf", filename=path.name)
        if msg.is_multipart():
            msg.set_boundary(f"==bench-{m['key']}==")  # deterministic MIME boundary
        out_path = DOCS / f"06_delay_notice_thread_{i}.eml"
        out_path.write_bytes(msg.as_bytes())
        out.append((m["key"], out_path, m["body"]))
    return out


def render_photo(out: Path, seed: int = 3) -> None:
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    import reportlab

    rng = random.Random(seed)
    W, H = 1600, 1200
    img = Image.new("RGB", (W, H), (120, 118, 112))
    d = ImageDraw.Draw(img)
    # floor: horizontal gradient with expansion-joint lines; back wall; a conveyor frame
    for y in range(H):
        t = y / H
        d.line([(0, y), (W, y)], fill=(int(70 + 90 * t), int(68 + 88 * t), int(62 + 80 * t)))
    d.rectangle([0, 0, W, 520], fill=(176, 182, 188))
    for x in range(0, W, 210):
        d.line([(x, 0), (x, 520)], fill=(150, 156, 162), width=3)
    for i in range(6):
        y = 560 + i * 110
        d.line([(0, y), (W, y)], fill=(90, 88, 82), width=2)
    for x in range(200, 1400, 160):  # conveyor legs
        d.rectangle([x, 300, x + 28, 700], fill=(60, 80, 130))
    d.rectangle([180, 280, 1420, 330], fill=(40, 60, 110))
    d.rectangle([180, 240, 1420, 282], fill=(200, 200, 205))
    fdir = Path(reportlab.__file__).parent / "fonts"
    big = ImageFont.truetype(str(fdir / "VeraBd.ttf"), 58)
    mid = ImageFont.truetype(str(fdir / "VeraBd.ttf"), 40)
    small = ImageFont.truetype(str(fdir / "Vera.ttf"), 30)
    # sign board
    d.rectangle([330, 700, 1270, 1080], fill=(250, 214, 40), outline=(20, 20, 20), width=8)
    lines = [
        ("BAY 3 — CONVEYOR LINE 2", big),
        ("INSTALLATION SUSPENDED", big),
        ("PENDING FLOOR RE-LEVEL (GRID C4–C7)", mid),
        ("TIDEWELL FABRICATION LTD — SITE SUPERVISOR C. REID", small),
    ]
    y = 730
    for text, font in lines:
        w = d.textlength(text, font=font)
        d.text(((W - w) / 2, y), text, fill=(15, 15, 15), font=font)
        y += font.size + 24
    # camera timestamp
    stamp = "2025-03-14 10:42"
    d.text((W - 420, H - 70), stamp, fill=(255, 170, 40), font=mid)
    # grain + slight blur to look like a phone photo
    img = img.filter(ImageFilter.GaussianBlur(0.8))
    noise = Image.effect_noise((W, H), 14).convert("RGB")
    img = Image.blend(img, noise, 0.06)
    img.save(str(out), "JPEG", quality=82)


# ----------------------------------------------------------------------------
# expected/
# ----------------------------------------------------------------------------



def org_name(bible: dict, ref: str) -> dict:
    for o in bible["organizations"]:
        if o["key"] == ref:
            return {"key": ref, "name": o["name"], "aliases": o.get("aliases", [])}
    p = person(bible, ref)
    return {"key": ref, "name": p["name"], "aliases": [p["name"].split()[-1]]}


def write_expected(bible: dict, docs_meta: list[dict], texts: dict[str, str]) -> None:
    EXPECTED.mkdir(exist_ok=True)
    (EXPECTED / "text").mkdir(exist_ok=True)
    for key, text in texts.items():
        (EXPECTED / "text" / f"{key}.txt").write_text(text)

    dump = lambda name, obj: (EXPECTED / name).write_text(yaml.safe_dump(obj, sort_keys=False, allow_unicode=True, width=120))

    dump("documents.yaml", {"case_version": bible["case_version"], "documents": docs_meta})
    dump(
        "chronology.yaml",
        {
            "case_version": bible["case_version"],
            "events": [
                {"key": e["key"], "date": e["date"], "text": e["text"], "docs": e["docs"]} for e in bible["events"]
            ],
        },
    )
    dump("families.yaml", {"case_version": bible["case_version"], "families": bible["families"], "copies": bible["copies"]})
    dump(
        "graph.yaml",
        {
            "case_version": bible["case_version"],
            "entities": [{"type": e["type"], **org_name(bible, e["ref"])} for e in bible["graph"]["entities"]],
            "relations": [
                {"from": org_name(bible, r["from"])["name"], "to": org_name(bible, r["to"])["name"], "kind": r["kind"]}
                for r in bible["graph"]["relations"]
            ],
        },
    )
    dump("search.yaml", {"case_version": bible["case_version"], "queries": bible["search"]})


def render_whiteboard(out: Path) -> None:
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    import reportlab

    fdir = Path(reportlab.__file__).parent / "fonts"
    W, H = 1400, 1000
    img = Image.new("RGB", (W, H), (236, 238, 236))
    d = ImageDraw.Draw(img)
    d.rectangle([40, 40, W - 40, H - 40], outline=(120, 120, 125), width=14)
    f1 = ImageFont.truetype(str(fdir / "VeraBd.ttf"), 54)
    f2 = ImageFont.truetype(str(fdir / "Vera.ttf"), 44)
    lines = [
        ("SITE MEETING 18/03/2025", f1, (20, 60, 160)),
        ("- 400A SUPPLY ENERGISED TODAY 18 MAR 2025 (utility + MP)", f2, (30, 30, 30)),
        ("- LOOP CHECKS BAY 1-2 START 19/03", f2, (30, 30, 30)),
        ("- BAY 3 C4-C7: SUB-FRAME OPTION, START 24/03", f2, (170, 30, 30)),
        ("- CO-3 STILL UNSIGNED -> PN TO CHASE LO", f2, (170, 30, 30)),
        ("- MECH COMPLETION NOTICE 21/03 (EXCL. C4-C7)", f2, (30, 30, 30)),
        ("TW / PN", f2, (20, 60, 160)),
    ]
    y = 110
    for text, font, col in lines:
        d.text((110, y), text, fill=col, font=font)
        y += font.size + 46
    img = img.rotate(-1.2, resample=Image.BICUBIC, expand=False, fillcolor=(200, 200, 200)).filter(ImageFilter.GaussianBlur(0.7))
    img.save(str(out), "PNG")



def render_agent_attachments() -> None:
    """Files attached in chat scenarios (agent/attachments/), from source/agent-attachments/."""
    src = SRC / "agent-attachments"
    out_dir = HERE / "agent" / "attachments"
    out_dir.mkdir(parents=True, exist_ok=True)
    for md in sorted(src.glob("*.md")):
        out = out_dir / (f"{md.stem}_2025-05-23.pdf" if md.stem == "brisamar_reply" else f"{md.stem}.pdf")
        render_pdf(parse_md(md.read_text(), SRC), out)
        canonicalize(out)
    render_whiteboard(out_dir / "whiteboard_2025-03-18.png")


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="overwrite docs/ even if present")
    args = ap.parse_args()

    bible = yaml.safe_load((HERE / "bible.yaml").read_text())
    if DOCS.exists() and any(DOCS.iterdir()) and not args.force:
        print("docs/ already populated; pass --force to regenerate (and bump case_version).", file=sys.stderr)
        return 2
    DOCS.mkdir(exist_ok=True)
    shutil.rmtree(DOCS)
    DOCS.mkdir()

    docs_by_key = {d["key"]: d for d in bible["documents"]}
    docs_meta: list[dict] = []
    texts: dict[str, str] = {}
    rendered: dict[str, Path] = {}

    def meta(d: dict, files: list[str], page_count: int | None, **extra) -> dict:
        m = {
            "key": d["key"],
            "files": files,
            "title": d.get("title"),
            "date": d.get("date"),
            "date_tolerance": d.get("date_tolerance", "exact"),
            "types": d.get("types", []),
            "language": d.get("language"),
            "page_count": page_count,
            "must_phrases": d.get("must_phrases", []),
            "must_facts": d.get("must_facts", []),
            "family": d.get("family"),
        }
        for k in ("ocr_similarity_min", "injection_marker", "attached_to", "vision_phrases"):
            if d.get(k) is not None:
                m[k] = d[k]
        m.update(extra)
        return m

    # survey first: it is attached to email 1
    order = ["survey-attachment"] + [d["key"] for d in bible["documents"] if d["key"] != "survey-attachment"]
    for key in order:
        d = docs_by_key[key]
        fmt = d["format"]
        if fmt == "pdf":
            blocks = parse_md((SRC / d["source"]).read_text(), SRC)
            out = (ATTACH_ONLY if d.get("attached_to") else DOCS) / d["file"]
            out.parent.mkdir(exist_ok=True)
            n = render_pdf(blocks, out)
            canonicalize(out)
            rendered[key] = out
            texts[key] = blocks_to_text(blocks)
            docs_meta.append(meta(d, [d["file"]], n, sha256=sha256(out)))
        elif fmt == "pdf-scan":
            blocks = parse_md((SRC / d["source"]).read_text(), SRC)
            tmp = DOCS / (d["file"] + ".text.tmp")
            render_pdf(blocks, tmp, font_size=11.5)
            out = DOCS / d["file"]
            n = render_scanned_pdf(tmp, out)
            tmp.unlink()
            canonicalize(out)
            rendered[key] = out
            texts[key] = blocks_to_text(blocks)
            docs_meta.append(meta(d, [d["file"]], n, sha256=sha256(out)))
        elif fmt == "docx":
            blocks = parse_md((SRC / d["source"]).read_text(), SRC)
            out = DOCS / d["file"]
            render_docx(blocks, out)
            canonicalize(out)
            rendered[key] = out
            texts[key] = blocks_to_text(blocks)
            docs_meta.append(meta(d, [d["file"]], None, sha256=sha256(out)))
        elif fmt == "txt":
            blocks = parse_md((SRC / d["source"]).read_text(), SRC)
            out = DOCS / d["file"]
            out.write_text(blocks_to_text(blocks))
            rendered[key] = out
            texts[key] = blocks_to_text(blocks)
            docs_meta.append(meta(d, [d["file"]], None, sha256=sha256(out)))
        elif fmt == "xlsx":
            out = DOCS / d["file"]
            render_xlsx(bible, out)
            canonicalize(out)
            rendered[key] = out
            docs_meta.append(meta(d, [d["file"]], None, sha256=sha256(out)))
        elif fmt == "jpg":
            out = DOCS / d["file"]
            render_photo(out)
            rendered[key] = out
            docs_meta.append(meta(d, [d["file"]], 1, sha256=sha256(out)))
        elif fmt == "eml-thread":
            emails = yaml.safe_load((SRC / d["source"]).read_text())
            att = {a["doc"]: rendered[a["doc"]] for a in bible["email_attachments"]}
            msgs = render_eml_thread(bible, emails, att)
            for (mkey, path, body), m in zip(msgs, emails["messages"]):
                texts[mkey] = body
                docs_meta.append(
                    {
                        "key": mkey,
                        "files": [path.name],
                        "thread": d["key"],
                        "title": d["title"],
                        "date": parsedate_to_datetime(m["date"]).date(),
                        "date_tolerance": "exact",
                        "types": d["types"],
                        "language": d["language"],
                        "page_count": None,
                        "must_phrases": m.get("must_phrases", []),
                        "must_facts": m.get("must_facts", []),
                        "family": {"group": "delay-thread", "role": "message"},
                        "attachments": m.get("attachments", []),
                        "sha256": sha256(path),
                    }
                )
        elif fmt == "csv":
            out = DOCS / d["file"]
            rows = [["Milestone", "Share of contract price", "Amount EUR", "Invoice no", "Invoice date", "Due date", "Status", "Paid on"]]
            for r in bible["payment_schedule"]:
                rows.append([r["milestone"], f'{r["share_pct"]}%', r["amount_eur"], r["invoice"], r["invoice_date"], r["due"], r["status"], r["paid"]])
            rows.append(["Total", "100%", bible["contract"]["price_v2_eur"], "", "", "", "", ""])
            rows.append(["Contract", "Supply and Installation Agreement 4 Nov 2024 as amended and restated 20 Jan 2025 (NWL-VAL-2024-07); payment 30 days from invoice (clause 3.3)", "", "", "", "", "", ""])
            texts[key] = render_csv(rows, out)
            rendered[key] = out
            docs_meta.append(meta(d, [d["file"]], None, sha256=sha256(out)))
        elif fmt == "png-scan":
            blocks = parse_md((SRC / d["source"]).read_text(), SRC)
            out = DOCS / d["file"]
            render_png_scan(blocks, out)
            rendered[key] = out
            texts[key] = blocks_to_text(blocks)
            docs_meta.append(meta(d, [d["file"]], 1, sha256=sha256(out)))
        elif fmt == "webp-photo":
            out = DOCS / d["file"]
            render_sign_photo(out, [(ln, sz) for ln, sz in d["sign_lines"]], d["stamp"], fmt="WEBP", palette=d.get("palette", "warehouse"), seed=9)
            rendered[key] = out
            docs_meta.append(meta(d, [d["file"]], 1, sha256=sha256(out)))
        elif fmt == "pdf-long":
            blocks = parse_md((SRC / d["source"]).read_text(), SRC)
            out = DOCS / d["file"]
            cfg = d["long"]
            sections = 12
            while True:
                full = expand_long_document(
                    blocks, sections=sections, start_section=cfg["start_section"], clauses_per_section=24,
                    title_pool=LONG_TITLES, subject_pool=LONG_SUBJECTS, requirement_pool=LONG_REQUIREMENTS, standard_pool=LONG_STANDARDS,
                    anchors=cfg.get("anchors"),
                )
                n = render_pdf(full, out)
                if n >= cfg["min_pages"] or sections > 60:
                    break
                sections += 4
            canonicalize(out)
            rendered[key] = out
            texts[key] = blocks_to_text(full)
            docs_meta.append(meta(d, [d["file"]], n, sha256=sha256(out)))
        elif fmt == "eml-html":
            spec = yaml.safe_load((SRC / d["source"]).read_text())
            out = DOCS / d["file"]
            render_eml(
                out, subject=spec["subject"], from_addr=addr(bible, spec["from"]["ref"]), to=[addr(bible, t["ref"]) for t in spec["to"]],
                cc=[addr(bible, t["ref"]) for t in spec.get("cc", [])], date_hdr=spec["date"], message_id=spec["message_id"],
                text=spec["text"], html=spec["html"], inline_images={"nwl-logo": logo_png("BRISAMAR")}, boundary=f"==bench-{key}==",
            )
            rendered[key] = out
            texts[key] = spec["text"]
            docs_meta.append(meta(d, [d["file"]], None, sha256=sha256(out), must_facts=spec.get("must_facts", [])))
        elif fmt == "copy-of":
            src = rendered[d["copy_of"]]
            out = DOCS / d["file"]
            shutil.copyfile(src, out)
            docs_meta.append({"key": key, "files": [d["file"]], "copy_of": d["copy_of"], "sha256": sha256(out)})
        else:
            raise SystemExit(f"unknown format {fmt} for {key}")

    # the survey is an email attachment: it is not uploaded on its own
    for m in docs_meta:
        if m["key"] == "survey-attachment":
            m["uploaded_directly"] = False

    for m in docs_meta:
        for fname in m["files"]:
            fp = (ATTACH_ONLY if m.get("attached_to") else DOCS) / fname
            if fp.exists():
                m["sha256"] = sha256(fp)

    write_expected(bible, docs_meta, texts)
    render_agent_attachments()

    total_pages = sum(m.get("page_count") or 0 for m in docs_meta)
    uploads = [f for m in docs_meta if m.get("uploaded_directly", True) and "copy_of" not in m or m.get("copy_of") for f in m["files"]]
    manifest = {
        "case_version": bible["case_version"],
        "case_name": bible["case"]["name"],
        "upload": sorted(f for m in docs_meta if m.get("uploaded_directly", True) for f in m["files"]),
        "pdf_image_pages": total_pages,
        "filename_to_key": {f: m["key"] for m in docs_meta for f in m["files"]},
    }
    (EXPECTED / "manifest.yaml").write_text(yaml.safe_dump(manifest, sort_keys=False, allow_unicode=True))
    print(f"rendered {len(list(DOCS.glob('*.*')))} upload files; PDF/image pages: {total_pages}")
    for m in docs_meta:
        print(f"  {m['key']:<20} {m['files'][0]:<40} pages={m.get('page_count')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
