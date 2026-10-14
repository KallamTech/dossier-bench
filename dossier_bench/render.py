"""Shared document renderers for benchmark cases: a markdown subset → PDF / scanned PDF / DOCX / PNG scan /
long PDF / CSV / photo / HTML email, all byte-deterministic (fixed fonts, dates, MIME boundaries, zip timestamps)."""

from __future__ import annotations

import csv
import hashlib
import io
import random
import re
from datetime import date
from email.message import EmailMessage
from pathlib import Path


PAGEBREAK = "<<<PAGEBREAK>>>"
INCLUDE_RE = re.compile(r"^<<<INCLUDE (\S+)>>>$")


def parse_md(text: str, src_dir: Path | None = None) -> list[tuple]:
    blocks: list[tuple] = []
    para: list[str] = []
    table: list[list[str]] = []
    bullets: list[str] = []

    def flush():
        nonlocal para, table, bullets
        if para:
            blocks.append(("p", " ".join(para)))
            para = []
        if table:
            blocks.append(("table", table))
            table = []
        if bullets:
            blocks.append(("ul", bullets))
            bullets = []

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush()
            continue
        m = INCLUDE_RE.match(line.strip())
        if m:
            flush()
            blocks.extend(parse_md(((src_dir or Path.cwd()) / m.group(1)).read_text(), src_dir))
            continue
        if line.strip() == PAGEBREAK:
            flush()
            blocks.append(("pagebreak",))
        elif line.startswith("# "):
            flush()
            blocks.append(("h1", line[2:]))
        elif line.startswith("## "):
            flush()
            blocks.append(("h2", line[3:]))
        elif line.startswith("### "):
            flush()
            blocks.append(("h3", line[4:]))
        elif line.startswith("|"):
            if para or bullets:
                flush()
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r"-+", c) for c in cells):
                continue
            table.append(cells)
        elif line.startswith("- "):
            if para or table:
                flush()
            bullets.append(line[2:])
        else:
            if table or bullets:
                flush()
            para.append(line.strip())
    flush()
    return blocks


def strip_md(s: str) -> str:
    return s.replace("**", "").replace("*", "")


def blocks_to_text(blocks: list[tuple]) -> str:
    out: list[str] = []
    for b in blocks:
        kind = b[0]
        if kind in ("h1", "h2", "h3", "p"):
            out.append(strip_md(b[1]))
            out.append("")
        elif kind == "ul":
            out.extend("- " + strip_md(i) for i in b[1])
            out.append("")
        elif kind == "table":
            for row in b[1]:
                out.append(" | ".join(strip_md(c) for c in row))
            out.append("")
        elif kind == "pagebreak":
            out.append("")
    return "\n".join(out).strip() + "\n"


def _fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import reportlab

    fdir = Path(reportlab.__file__).parent / "fonts"
    try:
        pdfmetrics.registerFont(TTFont("Body", str(fdir / "Vera.ttf")))
        pdfmetrics.registerFont(TTFont("BodyBold", str(fdir / "VeraBd.ttf")))
        pdfmetrics.registerFont(TTFont("BodyItalic", str(fdir / "VeraIt.ttf")))
        pdfmetrics.registerFontFamily("Body", normal="Body", bold="BodyBold", italic="BodyItalic", boldItalic="BodyBold")
        return "Body", "BodyBold"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


def md_inline(s: str) -> str:
    s = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*(?!\*)", r"<i>\1</i>", s)
    return s


def render_pdf(blocks: list[tuple], out: Path, *, font_size: float = 12) -> int:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    body, bold = _fonts()
    st_p = ParagraphStyle("p", fontName=body, fontSize=font_size, leading=font_size * 1.55, spaceAfter=6)
    st_h1 = ParagraphStyle("h1", fontName=bold, fontSize=16, leading=20, spaceAfter=10, alignment=TA_CENTER)
    st_h2 = ParagraphStyle("h2", fontName=bold, fontSize=12.5, leading=16, spaceBefore=10, spaceAfter=6)
    st_h3 = ParagraphStyle("h3", fontName=bold, fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
    st_cell = ParagraphStyle("cell", fontName=body, fontSize=font_size - 1, leading=(font_size - 1) * 1.35)
    st_cellb = ParagraphStyle("cellb", parent=st_cell, fontName=bold)

    story = []
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            story.append(Paragraph(md_inline(b[1]), st_h1))
        elif kind == "h2":
            story.append(Paragraph(md_inline(b[1]), st_h2))
        elif kind == "h3":
            story.append(Paragraph(md_inline(b[1]), st_h3))
        elif kind == "p":
            story.append(Paragraph(md_inline(b[1]), st_p))
        elif kind == "ul":
            story.append(ListFlowable([ListItem(Paragraph(md_inline(i), st_p)) for i in b[1]], bulletType="bullet"))
        elif kind == "table":
            rows = b[1]
            ncols = max(len(r) for r in rows)
            data = []
            for ri, r in enumerate(rows):
                r = r + [""] * (ncols - len(r))
                data.append([Paragraph(md_inline(c), st_cellb if ri == 0 else st_cell) for c in r])
            avail = A4[0] - 50 * mm
            t = Table(data, colWidths=[avail / ncols] * ncols, repeatRows=1)
            t.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8e8")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.append(t)
            story.append(Spacer(1, 6))
        elif kind == "pagebreak":
            story.append(PageBreak())

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(body, 8)
        canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        str(out), pagesize=A4, leftMargin=25 * mm, rightMargin=25 * mm, topMargin=25 * mm, bottomMargin=25 * mm,
        title=strip_md(next((b[1] for b in blocks if b[0] == "h1"), out.stem)), author="",
        invariant=1,  # fixed creation date + ids → byte-identical output for identical content
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return pdf_page_count(out)


def pdf_page_count(path: Path) -> int:
    from pypdf import PdfReader

    return len(PdfReader(str(path)).pages)


def render_scanned_pdf(text_pdf: Path, out: Path, seed: int = 7) -> int:
    """Rasterise a text PDF and re-save as image-only pages (no text layer)."""
    import pypdfium2 as pdfium
    from PIL import Image, ImageFilter

    rng = random.Random(seed)
    pdf = pdfium.PdfDocument(str(text_pdf))
    pages = []
    for i in range(len(pdf)):
        img = pdf[i].render(scale=150 / 72).to_pil().convert("L")
        img = img.rotate(rng.uniform(-0.6, 0.6), resample=Image.BICUBIC, expand=False, fillcolor=255)
        img = img.filter(ImageFilter.GaussianBlur(0.6))
        noise = Image.effect_noise(img.size, 18).convert("L")
        img = Image.blend(img, noise, 0.08)
        img = img.point(lambda v: 255 if v > 235 else int(v * 0.97))
        pages.append(img.convert("RGB"))
    pages[0].save(str(out), "PDF", save_all=True, append_images=pages[1:], resolution=150, quality=70)
    return len(pages)


def _docx_runs(par, text: str):
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if not chunk:
            continue
        run = par.add_run(chunk.replace("*", ""))
        run.bold = i % 2 == 1


def render_docx(blocks: list[tuple], out: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    from datetime import datetime

    doc.core_properties.created = doc.core_properties.modified = datetime(2025, 1, 20, 9, 0, 0)
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            doc.add_heading(strip_md(b[1]), level=1)
        elif kind == "h2":
            doc.add_heading(strip_md(b[1]), level=2)
        elif kind == "h3":
            doc.add_heading(strip_md(b[1]), level=3)
        elif kind == "p":
            _docx_runs(doc.add_paragraph(), b[1])
        elif kind == "ul":
            for item in b[1]:
                _docx_runs(doc.add_paragraph(style="List Bullet"), item)
        elif kind == "table":
            rows = b[1]
            ncols = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncols)
            t.style = "Table Grid"
            for ri, r in enumerate(rows):
                for ci in range(ncols):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    _docx_runs(cell.paragraphs[0], (r[ci] if ci < len(r) else "") if ri else f"**{strip_md(r[ci]) if ci < len(r) else ''}**")
            doc.add_paragraph()
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(str(out))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def canonicalize(path: Path) -> None:
    """Strip the last sources of nondeterminism so identical content gives identical bytes:
    zip entry timestamps in docx/xlsx and Pillow's PDF CreationDate/ModDate (same-length replace keeps xref offsets)."""
    import zipfile

    if path.suffix.lower() in (".docx", ".xlsx"):
        tmp = path.with_suffix(path.suffix + ".tmp")
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in sorted(zin.infolist(), key=lambda i: i.filename):
                data = zin.read(info.filename)
                if info.filename == "docProps/core.xml":  # openpyxl stamps dcterms:modified at save time
                    data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>2025-01-01T00:00:00Z", data)
                zi = zipfile.ZipInfo(info.filename, date_time=(2025, 1, 1, 0, 0, 0))
                zi.compress_type = zipfile.ZIP_DEFLATED
                zout.writestr(zi, data)
        tmp.replace(path)
    elif path.suffix.lower() == ".pdf":
        b = path.read_bytes()
        fixed = re.sub(rb"(/(?:CreationDate|ModDate)\s*\(D:)\d{14}", lambda m: m.group(1) + b"20250101000000", b)
        if fixed != b:
            path.write_bytes(fixed)


# ----------------------------------------------------------------------------
# additional formats (2026-10-13): PNG scan, long PDF, CSV, generic photo, HTML email
# ----------------------------------------------------------------------------


def render_png_scan(blocks: list[tuple], out: Path, seed: int = 11) -> None:
    """One page of text rendered as a phone scan: PNG, no text layer, slight skew + noise."""
    import pypdfium2 as pdfium
    from PIL import Image, ImageFilter

    tmp = out.with_suffix(".tmp.pdf")
    render_pdf(blocks, tmp, font_size=11)
    rng = random.Random(seed)
    pdf = pdfium.PdfDocument(str(tmp))
    img = pdf[0].render(scale=160 / 72).to_pil().convert("L")
    pdf.close()
    tmp.unlink()
    img = img.rotate(rng.uniform(-1.5, 1.5), resample=Image.BICUBIC, expand=False, fillcolor=255)
    img = img.filter(ImageFilter.GaussianBlur(0.5))
    noise = Image.effect_noise(img.size, 22).convert("L")
    img = Image.blend(img, noise, 0.10)
    img = img.point(lambda v: 255 if v > 225 else int(v * 0.93))
    img.convert("RGB").save(str(out), "PNG", optimize=True)


def expand_long_document(blocks: list[tuple], *, sections: int, start_section: int, clauses_per_section: int, title_pool: list[str], subject_pool: list[str], requirement_pool: list[str], standard_pool: list[str], anchors: dict[str, str] | None = None) -> list[tuple]:
    """Append deterministic numbered sections (headings, clauses, a table per section) to `blocks`, so a
    specification can run to dozens of pages. `anchors` = {"12.7": "text"} pins specific clause text."""
    anchors = anchors or {}
    out = list(blocks)
    for n in range(sections):
        s = start_section + n
        out.append(("h2", f"{s}. {title_pool[n % len(title_pool)]}"))
        for i in range(1, clauses_per_section + 1):
            key = f"{s}.{i}"
            if key in anchors:
                out.append(("p", f"{key} {anchors[key]}"))
                continue
            subj = subject_pool[(s * 7 + i) % len(subject_pool)]
            req = requirement_pool[(s * 3 + i) % len(requirement_pool)]
            std = standard_pool[(s + i) % len(standard_pool)]
            out.append(("p", f"{key} The {subj} shall {req} in accordance with {std}. Compliance shall be demonstrated at factory acceptance testing and again at site acceptance testing, and recorded in the commissioning dossier."))
        rows = [[f"{s}.T{r}", subject_pool[(s + r) % len(subject_pool)], standard_pool[(s * r) % len(standard_pool)], f"{(s * 13 + r * 7) % 90 + 10} days"] for r in range(1, 7)]
        out.append(("table", [["Ref", "Item", "Standard", "Lead time"]] + rows))
    return out


def render_csv(rows: list[list], out: Path) -> str:
    """Write rows as CSV (CRLF, UTF-8) and return the same content as plain text for expected/text."""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    for r in rows:
        w.writerow(r)
    out.write_text(buf.getvalue(), encoding="utf-8")
    return "\n".join(", ".join(str(c) for c in r) for r in rows)


def render_sign_photo(out: Path, lines: list[tuple[str, str]], stamp: str, *, fmt: str = "JPEG", seed: int = 3, palette: str = "warehouse") -> None:
    """A phone photo of a sign board (or a measurement card) on a site; `lines` = [(text, size)] with
    size in {big, mid, small}; `fmt` JPEG or WEBP."""
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    import reportlab

    rng = random.Random(seed)
    W, H = 1600, 1200
    img = Image.new("RGB", (W, H), (120, 118, 112))
    d = ImageDraw.Draw(img)
    for y in range(H):
        t = y / H
        d.line([(0, y), (W, y)], fill=(int(70 + 90 * t), int(68 + 88 * t), int(62 + 80 * t)))
    if palette == "warehouse":
        d.rectangle([0, 0, W, 520], fill=(176, 182, 188))
        for x in range(0, W, 210):
            d.line([(x, 0), (x, 520)], fill=(150, 156, 162), width=3)
    else:  # bare slab: grid chalk lines
        d.rectangle([0, 0, W, H], fill=(150, 148, 142))
        for x in range(0, W, 230):
            d.line([(x, 0), (x, H)], fill=(205, 205, 200), width=4)
        for y in range(0, H, 230):
            d.line([(0, y), (W, y)], fill=(205, 205, 200), width=4)
    fdir = Path(reportlab.__file__).parent / "fonts"
    fonts = {"big": ImageFont.truetype(str(fdir / "VeraBd.ttf"), 58), "mid": ImageFont.truetype(str(fdir / "VeraBd.ttf"), 40), "small": ImageFont.truetype(str(fdir / "Vera.ttf"), 30)}
    d.rectangle([330, 700, 1270, 1080], fill=(250, 214, 40) if palette == "warehouse" else (245, 245, 240), outline=(20, 20, 20), width=8)
    y = 730
    for text, size in lines:
        f = fonts[size]
        w = d.textlength(text, font=f)
        d.text(((W - w) / 2, y), text, fill=(15, 15, 15), font=f)
        y += f.size + 24
    d.text((W - 420, H - 70), stamp, fill=(255, 170, 40), font=fonts["mid"])
    img = img.filter(ImageFilter.GaussianBlur(0.8))
    noise = Image.effect_noise((W, H), 14).convert("RGB")
    img = Image.blend(img, noise, 0.06)
    if fmt.upper() == "WEBP":
        img.save(str(out), "WEBP", quality=80, method=4)
    else:
        img.save(str(out), "JPEG", quality=82)


def logo_png(text: str, seed: int = 5) -> bytes:
    """A small deterministic PNG word-mark, for inline email images."""
    from PIL import Image, ImageDraw, ImageFont
    import reportlab

    fdir = Path(reportlab.__file__).parent / "fonts"
    img = Image.new("RGB", (320, 80), (14, 60, 120))
    d = ImageDraw.Draw(img)
    d.text((14, 20), text, fill=(255, 255, 255), font=ImageFont.truetype(str(fdir / "VeraBd.ttf"), 32))
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=True)
    return buf.getvalue()


def render_eml(out: Path, *, subject: str, from_addr: str, to: list[str], cc: list[str], date_hdr: str, message_id: str, text: str, html: str | None = None, inline_images: dict[str, bytes] | None = None, attachments: list[Path] | None = None, boundary: str = "==bench-eml==") -> None:
    """A single RFC822 message; with `html` the body is multipart/alternative, inline images are related by cid."""
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = from_addr
    msg["To"] = ", ".join(to)
    if cc:
        msg["Cc"] = ", ".join(cc)
    msg["Date"] = date_hdr
    msg["Message-ID"] = message_id
    msg.set_content(text.rstrip() + "\n")
    if html:
        msg.add_alternative(html, subtype="html")
        for cid, data in (inline_images or {}).items():
            msg.get_payload()[1].add_related(data, maintype="image", subtype="png", cid=f"<{cid}>", filename=f"{cid}.png")
    for p in attachments or []:
        msg.add_attachment(p.read_bytes(), maintype="application", subtype="octet-stream", filename=p.name)
    if msg.is_multipart():
        msg.set_boundary(boundary)
        for i, part in enumerate(msg.iter_parts()):
            if part.is_multipart():
                part.set_boundary(f"{boundary}{i}==")
    out.write_bytes(msg.as_bytes())


# ----------------------------------------------------------------------------
# generic spreadsheet + email thread (second case onwards)
# ----------------------------------------------------------------------------


def render_xlsx_spec(sheets: list[dict], out: Path) -> str:
    """`sheets` = [{"title", "rows": [[...], ...], "widths": [..], "bold_rows": [idx, ...]}]; returns the
    text a reader would see, for expected/text."""
    from datetime import datetime

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    lines: list[str] = []
    for n, sh in enumerate(sheets):
        ws = wb.active if n == 0 else wb.create_sheet()
        ws.title = sh["title"]
        lines.append(f"[{sh['title']}]")
        for r_i, row in enumerate(sh["rows"]):
            ws.append(row)
            if r_i in sh.get("bold_rows", []):
                for c in ws[r_i + 1]:
                    c.font = Font(bold=True)
            lines.append(", ".join("" if v is None else str(v) for v in row))
        for col, w in zip("ABCDEFGHIJKL", sh.get("widths", [])):
            ws.column_dimensions[col].width = w
    wb.properties.created = wb.properties.modified = datetime(2025, 1, 1, 0, 0, 0)
    wb.save(str(out))
    return "\n".join(lines)


def render_eml_thread(messages: list[dict], *, thread_subject: str, out_dir: Path, file_prefix: str, addr_of, name_of, attachments: dict[str, Path] | None = None) -> list[tuple[str, Path, str]]:
    """Render a reply chain to `<file_prefix>_{n}.eml`; replies quote the parent. `addr_of(ref)` →
    '"Name" <email>', `name_of(ref)` → 'Name'. Returns [(message key, path, plain body)]."""
    from email.utils import parsedate_to_datetime

    attachments = attachments or {}
    out: list[tuple[str, Path, str]] = []
    by_key = {m["key"]: m for m in messages}
    rendered: dict[str, str] = {}
    for i, m in enumerate(messages, start=1):
        msg = EmailMessage()
        msg["Subject"] = thread_subject if not m.get("in_reply_to") else "RE: " + thread_subject
        msg["From"] = addr_of(m["from"]["ref"])
        msg["To"] = ", ".join(addr_of(t["ref"]) for t in m["to"])
        if m.get("cc"):
            msg["Cc"] = ", ".join(addr_of(t["ref"]) for t in m["cc"])
        msg["Date"] = m["date"]
        msg["Message-ID"] = m["message_id"]
        body = m["body"].rstrip() + "\n"
        if m.get("in_reply_to"):
            parent = by_key[m["in_reply_to"]]
            msg["In-Reply-To"] = parent["message_id"]
            refs, p = [], parent
            while p is not None:
                refs.insert(0, p["message_id"])
                p = by_key.get(p.get("in_reply_to")) if p.get("in_reply_to") else None
            msg["References"] = " ".join(refs)
            pdt = parsedate_to_datetime(parent["date"])
            quoted = "\n".join("> " + ln for ln in rendered[parent["key"]].splitlines())
            body += f"\n\nOn {pdt.strftime('%a, %d %b %Y at %H:%M')}, {name_of(parent['from']['ref'])} wrote:\n\n{quoted}\n"
        rendered[m["key"]] = body
        msg.set_content(body)
        for att_key in m.get("attachments", []):
            path = attachments[att_key]
            msg.add_attachment(path.read_bytes(), maintype="application", subtype="pdf", filename=path.name)
        if msg.is_multipart():
            msg.set_boundary(f"==bench-{m['key']}==")
        path = out_dir / f"{file_prefix}_{i}.eml"
        path.write_bytes(msg.as_bytes())
        out.append((m["key"], path, m["body"]))
    return out
