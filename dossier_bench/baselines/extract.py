"""Plain text out of the corpus formats. No OCR here: scans and photos come back as images for the caller."""

from __future__ import annotations

import csv
import email
import io
import re
from dataclasses import dataclass, field
from email import policy
from pathlib import Path

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".gif"}


@dataclass
class Extracted:
    filename: str
    text: str = ""
    page_count: int | None = None
    images: list[tuple[str, bytes]] = field(default_factory=list)   # (mime, bytes) needing vision
    attachments: list[tuple[str, bytes]] = field(default_factory=list)  # (filename, bytes) inside an email


def strip_html(html: str) -> str:
    html = re.sub(r"(?is)<(script|style).*?</\1>", " ", html)
    html = re.sub(r"(?i)<br\s*/?>|</p>|</tr>|</div>|</h\d>", "\n", html)
    text = re.sub(r"<[^>]+>", " ", html)
    for a, b in (("&nbsp;", " "), ("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">"), ("&quot;", '"'), ("&#39;", "'")):
        text = text.replace(a, b)
    return re.sub(r"[ \t]+", " ", re.sub(r"\n\s*\n+", "\n\n", text)).strip()


def _pdf(data: bytes, ex: Extracted, render_scans: bool) -> None:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    ex.page_count = len(reader.pages)
    pages = [(p.extract_text() or "").strip() for p in reader.pages]
    ex.text = "\n\n".join(pages)
    if render_scans and len(ex.text) < 50 * max(1, ex.page_count):
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(data)
        for i in range(min(len(pdf), 10)):
            buf = io.BytesIO()
            pdf[i].render(scale=1.5).to_pil().save(buf, format="PNG")
            ex.images.append(("image/png", buf.getvalue()))


def _docx(data: bytes, ex: Extracted) -> None:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += ["\t".join(c.text for c in row.cells) for row in t.rows]
    ex.text = "\n".join(parts)


def _xlsx(data: bytes, ex: Extracted) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                parts.append("\t".join("" if v is None else str(v) for v in row))
    ex.text = "\n".join(parts)


def _eml(data: bytes, ex: Extracted) -> None:
    msg = email.message_from_bytes(data, policy=policy.default)
    head = "\n".join(f"{h}: {msg[h]}" for h in ("Subject", "From", "To", "Cc", "Date") if msg[h])
    body = msg.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        text = body.get_content()
        if body.get_content_type() == "text/html":
            text = strip_html(text)
    ex.text = f"{head}\n\n{text}"
    for part in msg.iter_attachments():
        name = part.get_filename()
        if name and part.get_content_disposition() == "attachment":
            ex.attachments.append((name, part.get_payload(decode=True)))


def extract(filename: str, data: bytes, *, render_scans: bool = False) -> Extracted:
    ex = Extracted(filename)
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        _pdf(data, ex, render_scans)
    elif ext == ".docx":
        _docx(data, ex)
    elif ext == ".xlsx":
        _xlsx(data, ex)
    elif ext == ".eml":
        _eml(data, ex)
    elif ext in (".html", ".htm"):
        ex.text = strip_html(data.decode("utf-8", "replace"))
    elif ext == ".csv":
        ex.text = "\n".join("\t".join(r) for r in csv.reader(io.StringIO(data.decode("utf-8-sig", "replace"))))
    elif ext in IMAGE_EXT:
        mime = "image/jpeg" if ext in (".jpg", ".jpeg") else f"image/{ext[1:]}"
        ex.images.append((mime, data))
    else:
        ex.text = data.decode("utf-8", "replace")
    return ex
